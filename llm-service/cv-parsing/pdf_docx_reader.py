"""
pdf_docx_reader.py — raw text extraction from uploaded CV files.

Per Section 2.3 of the sprint plan, the extraction library choice is yours
alone; it doesn't affect any other lane. This module supports PDF and DOCX
(the two formats candidates realistically upload) and fails loudly with a
typed exception rather than silently returning empty text, so extractor.py
can decide how to handle it (fallback path).
"""
from __future__ import annotations

import io
from typing import Optional


class UnsupportedFileTypeError(ValueError):
    """Raised when the file is neither a PDF nor a DOCX (by content or name)."""


class TextExtractionError(RuntimeError):
    """Raised when the file's type is supported but text extraction still fails
    (corrupt file, scanned/image-only PDF with no text layer, etc.)."""


def _looks_like_pdf(file_bytes: bytes) -> bool:
    return file_bytes[:5] == b"%PDF-"


def _looks_like_docx(file_bytes: bytes) -> bool:
    # DOCX is a ZIP archive; ZIP files start with "PK\x03\x04".
    return file_bytes[:4] == b"PK\x03\x04"


def _detect_file_type(file_bytes: bytes, filename: Optional[str] = None) -> str:
    """Returns 'pdf' or 'docx'. Prefers content sniffing over the filename
    extension, since uploaded filenames can't always be trusted, but falls
    back to the extension if the magic bytes are ambiguous."""
    if not file_bytes:
        raise TextExtractionError("Received an empty file (0 bytes).")

    if _looks_like_pdf(file_bytes):
        return "pdf"
    if _looks_like_docx(file_bytes):
        return "docx"

    if filename:
        lower = filename.lower()
        if lower.endswith(".pdf"):
            return "pdf"
        if lower.endswith(".docx"):
            return "docx"

    raise UnsupportedFileTypeError(
        "Could not determine file type from content or filename. "
        "Only PDF and DOCX are supported."
    )


def _extract_pdf_text(file_bytes: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise RuntimeError(
            "pypdf is not installed. Run: pip install pypdf"
        ) from e

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page in reader.pages:
            pages_text.append(page.extract_text() or "")
        text = "\n".join(pages_text).strip()
    except Exception as e:
        raise TextExtractionError(f"Failed to parse PDF: {e}") from e

    if not text:
        raise TextExtractionError(
            "PDF parsed successfully but contained no extractable text "
            "(likely a scanned/image-only PDF with no text layer)."
        )
    return text


def _extract_docx_text(file_bytes: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:
        raise RuntimeError(
            "python-docx is not installed. Run: pip install python-docx"
        ) from e

    try:
        document = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs if p.text.strip()]

        # Tables (skills matrices, experience tables) are common in CVs —
        # don't silently drop them.
        for table in document.tables:
            for row in table.rows:
                cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells_text:
                    parts.append(" | ".join(cells_text))

        text = "\n".join(parts).strip()
    except Exception as e:
        raise TextExtractionError(f"Failed to parse DOCX: {e}") from e

    if not text:
        raise TextExtractionError("DOCX parsed successfully but contained no text.")
    return text


def extract_text(file_bytes: bytes, filename: Optional[str] = None) -> str:
    """Extract raw text from a CV file (PDF or DOCX).

    Args:
        file_bytes: raw bytes of the uploaded file.
        filename: optional original filename, used as a fallback signal
            for type detection.

    Returns:
        Extracted raw text as a single string.

    Raises:
        UnsupportedFileTypeError: file is neither PDF nor DOCX.
        TextExtractionError: file type is supported but text could not
            be extracted (corrupt file, no text layer, etc.).
    """
    file_type = _detect_file_type(file_bytes, filename)
    if file_type == "pdf":
        return _extract_pdf_text(file_bytes)
    return _extract_docx_text(file_bytes)
